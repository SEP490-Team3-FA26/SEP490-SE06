import sys
import os
import re

if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_styled_paragraph(doc, text, style_name=None, space_before=0, space_after=6, line_spacing=1.15):
    p = doc.add_paragraph(style=style_name)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    
    # Process inline formatting: bold, italic, code
    # Regex splits while keeping tokens
    pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))'
    tokens = re.split(pattern, text)
    
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = p.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = p.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(199, 37, 78)
        elif token.startswith('[') and ']' in token and '(' in token and token.endswith(')'):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', token)
            if m:
                link_text = m.group(1)
                run = p.add_run(link_text)
                run.font.color.rgb = RGBColor(37, 99, 235)
                run.underline = True
            else:
                p.add_run(token)
        else:
            p.add_run(token)
    return p

def convert_md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    doc = Document()
    
    # Page setup - A4, 1 inch margins
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Default font styling
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(33, 37, 41)
    
    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        stripped = line.strip()
        
        # Code block handling
        if stripped.startswith('```'):
            if in_code_block:
                # End code block
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F8F9FA")
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                
                # set single border
                tcPr = cell._tc.get_or_add_tcPr()
                borders = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>'
                    f'<w:top w:val="single" w:sz="4" w:color="E2E8F0"/>'
                    f'<w:left w:val="single" w:sz="12" w:color="3B82F6"/>'
                    f'<w:bottom w:val="single" w:sz="4" w:color="E2E8F0"/>'
                    f'<w:right w:val="single" w:sz="4" w:color="E2E8F0"/>'
                    f'</w:tcBorders>'
                )
                tcPr.append(borders)
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                run = p.add_run("\n".join(code_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
                
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # Table handling
        if stripped.startswith('|') and stripped.endswith('|'):
            # Check if it's separator row: |:---|:---|
            if re.match(r'^\|[\s:-]+\|', stripped) and all(c in '|:- \t' for c in stripped):
                i += 1
                continue
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            table_rows.append(cells)
            i += 1
            continue
        else:
            if table_rows:
                # Flush table
                cols_count = max(len(r) for r in table_rows)
                tbl = doc.add_table(rows=len(table_rows), cols=cols_count)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl, color="CBD5E1", sz="4")
                
                for r_idx, row_data in enumerate(table_rows):
                    is_header = (r_idx == 0)
                    row = tbl.rows[r_idx]
                    for c_idx in range(cols_count):
                        val = row_data[c_idx] if c_idx < len(row_data) else ""
                        cell = row.cells[c_idx]
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                        
                        if is_header:
                            set_cell_background(cell, "1E3A8A") # Navy Blue
                        elif r_idx % 2 == 1:
                            set_cell_background(cell, "FFFFFF")
                        else:
                            set_cell_background(cell, "F8FAFC")
                            
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        p.paragraph_format.line_spacing = 1.1
                        
                        # Process cell formatting
                        pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|<br\s*/?>)'
                        tokens = re.split(pattern, val)
                        for token in tokens:
                            if not token:
                                continue
                            if re.match(r'<br\s*/?>', token):
                                p.add_run("\n")
                            elif token.startswith('**') and token.endswith('**'):
                                run = p.add_run(token[2:-2])
                                run.bold = True
                                if is_header:
                                    run.font.color.rgb = RGBColor(255, 255, 255)
                            elif token.startswith('*') and token.endswith('*'):
                                run = p.add_run(token[1:-1])
                                run.italic = True
                                if is_header:
                                    run.font.color.rgb = RGBColor(255, 255, 255)
                            elif token.startswith('`') and token.endswith('`'):
                                run = p.add_run(token[1:-1])
                                run.font.name = 'Consolas'
                                run.font.size = Pt(9)
                                run.font.color.rgb = RGBColor(220, 38, 38) if not is_header else RGBColor(254, 240, 138)
                            else:
                                run = p.add_run(token)
                                if is_header:
                                    run.bold = True
                                    run.font.color.rgb = RGBColor(255, 255, 255)
                table_rows = []
                doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # Empty line
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if stripped in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        # Callout alerts / Blockquotes
        if stripped.startswith('>'):
            quote_text = stripped.lstrip('>').strip()
            
            # Check alert type
            border_color = "0284C7" # Light Blue
            bg_color = "F0F9FF"
            title = ""
            
            if quote_text.startswith('[!NOTE]'):
                border_color = "3B82F6"
                bg_color = "EFF6FF"
                quote_text = quote_text.replace('[!NOTE]', '').strip()
                title = "📌 LƯU Ý:"
            elif quote_text.startswith('[!IMPORTANT]'):
                border_color = "8B5CF6"
                bg_color = "F5F3FF"
                quote_text = quote_text.replace('[!IMPORTANT]', '').strip()
                title = "⭐ QUAN TRỌNG:"
            elif quote_text.startswith('[!WARNING]') or quote_text.startswith('[!CAUTION]'):
                border_color = "EF4444"
                bg_color = "FEF2F2"
                quote_text = re.sub(r'\[!(WARNING|CAUTION)\]', '', quote_text).strip()
                title = "⚠️ CẢNH BÁO:"
            elif quote_text.startswith('[!TIP]'):
                border_color = "10B981"
                bg_color = "ECFDF5"
                quote_text = quote_text.replace('[!TIP]', '').strip()
                title = "💡 MẸO:"

            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:left w:val="single" w:sz="18" w:color="{border_color}"/>'
                f'<w:top w:val="none"/>'
                f'<w:bottom w:val="none"/>'
                f'<w:right w:val="none"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(borders)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            if title:
                r_title = p.add_run(f"{title} ")
                r_title.bold = True
                r_title.font.color.rgb = RGBColor(30, 41, 59)
                
            # Inline formatting for quote
            tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', quote_text)
            for token in tokens:
                if not token:
                    continue
                if token.startswith('**') and token.endswith('**'):
                    run = p.add_run(token[2:-2])
                    run.bold = True
                elif token.startswith('*') and token.endswith('*'):
                    run = p.add_run(token[1:-1])
                    run.italic = True
                elif token.startswith('`') and token.endswith('`'):
                    run = p.add_run(token[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9.5)
                else:
                    p.add_run(token)
                    
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # Headings
        if stripped.startswith('# '):
            p = add_styled_paragraph(doc, stripped[2:], space_before=12, space_after=6)
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(20)
                r.font.color.rgb = RGBColor(30, 58, 138) # Deep Navy
            i += 1
            continue
        elif stripped.startswith('## '):
            p = add_styled_paragraph(doc, stripped[3:], space_before=12, space_after=4)
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(15)
                r.font.color.rgb = RGBColor(30, 58, 138)
            i += 1
            continue
        elif stripped.startswith('### '):
            p = add_styled_paragraph(doc, stripped[4:], space_before=8, space_after=3)
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(13)
                r.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue
        elif stripped.startswith('#### '):
            p = add_styled_paragraph(doc, stripped[5:], space_before=6, space_after=2)
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(11.5)
                r.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        # Bullet lists
        if stripped.startswith('* ') or stripped.startswith('- '):
            indent_level = (len(line) - len(line.lstrip())) // 2
            text = stripped[2:]
            p = add_styled_paragraph(doc, text, space_before=1, space_after=2)
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            run_bullet = p.runs[0] if p.runs else p.add_run()
            # Add bullet symbol
            p_bullet = doc.paragraphs[-1]
            # Prefix bullet
            p.runs[0].text = "•  " + p.runs[0].text if p.runs else "•  "
            i += 1
            continue

        # Numbered lists
        m_num = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if m_num:
            num = m_num.group(1)
            text = m_num.group(2)
            indent_level = (len(line) - len(line.lstrip())) // 2
            p = add_styled_paragraph(doc, text, space_before=1, space_after=2)
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            if p.runs:
                p.runs[0].text = f"{num}.  " + p.runs[0].text
            else:
                p.add_run(f"{num}.  ")
            i += 1
            continue

        # Normal paragraph
        add_styled_paragraph(doc, stripped, space_before=2, space_after=5)
        i += 1

    # Overwrite directly to target docx
    doc.save(docx_path)
    print(f"✅ Document successfully saved/overwritten to: {docx_path}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python md_to_docx.py <input.md> [output.docx]")
        sys.exit(1)
    input_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(input_file)[0] + ".docx"
    convert_md_to_docx(input_file, output_file)
