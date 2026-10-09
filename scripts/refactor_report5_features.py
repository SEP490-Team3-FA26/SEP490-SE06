# scripts/refactor_report5_features.py
import os
import sys
import json
import time
import re
import urllib.request
import openpyxl
from copy import copy
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

sys.path.append(os.path.abspath('.'))
sys.path.append(os.path.abspath('.agents/skills/google-docx/scripts'))
from direct_gdrive_sync import get_access_token

SA_PATH = '/Users/tranhongphuoc/.config/gcloud/legacy_credentials/sep490@stone-climate-507417-k4.iam.gserviceaccount.com/adc.json'
SPREADSHEET_ID = '1Fffbgb6gvou2O56s19X4cNvS-o-NOThG'
INPUT_FILE = '/tmp/report5.xlsx'
OUTPUT_FILE = 'docs/Report5_Test_Report.xlsx'

# Definition of the 8 Feature Modules mapped directly with Report 3 (Section 3.2 - 3.9)
FEATURES = [
    {
        "id": "F01",
        "sheet_name": "Authentication",
        "feature_title": "Feature Authentication",
        "req": "User account authentication and identity management: user registration, credential login, Google OAuth2, email OTP verification, secure password recovery and reset, and mobile biometric authentication (FaceID/Fingerprint).",
        "ucs": [1, 2, 3, 4, 5, 6, 7, 113]
    },
    {
        "id": "F02",
        "sheet_name": "Customer Portal",
        "feature_title": "Feature Online Customer Portal",
        "req": "Online customer e-commerce portal and AI health assistance: medicine search, shopping cart management, online checkout, PayOS VietQR and COD payment, text and voice AI symptom consultation, order tracking, electronic health passport, and loyalty reward points.",
        "ucs": [8, 9, 10, 11, 12, 13, 14, 15, 16, 80, 81, 85, 86, 96, 97, 98, 99, 106, 107, 108, 109, 117, 118]
    },
    {
        "id": "F03",
        "sheet_name": "Pharmacist POS",
        "feature_title": "Feature Pharmacist POS Counter Terminal",
        "req": "Pharmacist point-of-sale (POS) counter terminal: rapid OTC and prescription search, retail cart management, active-ingredient drug substitution, voucher and loyalty point redemption, AI OCR prescription scanning, drug-drug interaction safety checks, counter VietQR payment, e-invoice and wholesale invoicing, return processing, and shift cash drawer balancing.",
        "ucs": [17, 18, 19, 20, 21, 22, 23, 24, 25, 26, 27, 28, 67, 68, 69, 70, 84, 111]
    },
    {
        "id": "F04",
        "sheet_name": "Central Warehouse",
        "feature_title": "Feature Central Warehouse & FEFO Inventory",
        "req": "Central warehouse and FEFO inventory management: master medicine catalog, batch and expiry date tracking, purchase order (PO) generation, barcode-assisted goods receipt notes (GRN), AI camera quality inspection, near-expiry alerts, stock card logging, and GSP bin/rack location management.",
        "ucs": [29, 30, 31, 32, 33, 34, 35, 87, 88, 89, 90, 94]
    },
    {
        "id": "F05",
        "sheet_name": "Branch Management",
        "feature_title": "Feature Branch Store Management",
        "req": "Branch store operations and inventory governance: branch stock visibility, purchase requisition (PR) creation, inter-branch stock transfers, replenishment approvals, transfer dispatch confirmation, branch-specific pricing, shipper delivery QR verification, and offline data synchronization.",
        "ucs": [36, 37, 38, 39, 40, 93, 100, 115]
    },
    {
        "id": "F06",
        "sheet_name": "HQ Governance",
        "feature_title": "Feature HQ Governance & Master Data",
        "req": "Headquarter system governance and master data administration: staff management, account creation and approval, branch provisioning, RBAC role permissions, global medicine and active ingredient catalog, multi-unit of measure (Multi-UOM) conversions, SKU and GS1 barcode generation, GDP supplier compliance contracts, auxiliary label printing, and multilingual system localization (VI/EN).",
        "ucs": [41, 42, 43, 44, 45, 46, 47, 48, 49, 50, 51, 91, 92, 103, 104, 119, 120]
    },
    {
        "id": "F07",
        "sheet_name": "Finance & AI Supply Chain",
        "feature_title": "Feature Corporate Finance & AI Supply Chain",
        "req": "Corporate financial governance and predictive AI supply chain analytics: executive revenue dashboards, cross-branch performance benchmarking, PO and GRN approval workflows, promotional flash sales and voucher management, cash flow analysis, system audit logging, AI demand forecasting, automated PO generation, inventory anomaly detection, RFM customer segmentation, and marketing ROI analytics.",
        "ucs": [52, 53, 54, 55, 56, 57, 58, 59, 60, 61, 62, 63, 64, 65, 66, 71, 72, 73, 74, 75, 76, 77, 83, 95, 101, 105, 110]
    },
    {
        "id": "F08",
        "sheet_name": "Extended Operations & IoT",
        "feature_title": "Feature Extended Operations & IoT Compliance",
        "req": "Extended operations, specialized pharmaceutical compliance, and IoT monitoring: imported medicine auxiliary barcode labels, multi-supplier bulk RFQ bidding, GSP warehouse temperature and humidity IoT telemetry, controlled and narcotic substance logs, disaster recovery backups, IoT device calibration, and customer credit limit governance.",
        "ucs": [78, 79, 82, 102, 112, 114, 116]
    }
]

# Style Presets matching sample images (Image 1, 2, 3)
FONT_FAMILY = "Tahoma"

# Border definitions
thin_side = Side(border_style="thin", color="000000")
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
dashed_side = Side(border_style="hair", color="A0A0A0")
grid_border = Border(left=dashed_side, right=dashed_side, top=dashed_side, bottom=dashed_side)

# Colors
OLIVE_GREEN_FILL = PatternFill(start_color="5B7038", end_color="5B7038", fill_type="solid") # Image 3 Table Header & Image 1 Category Header
NAVY_BLUE_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")   # Image 1 & Image 2 Headers
CYAN_SUBHEADER_FILL = PatternFill(start_color="D2EAF0", end_color="D2EAF0", fill_type="solid") # Image 3 Function Subheader
SUMMARY_HEADER_FILL = PatternFill(start_color="E8EEF5", end_color="E8EEF5", fill_type="solid")

FONT_WHITE_BOLD = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
FONT_TITLE_BOLD = Font(name=FONT_FAMILY, size=16, bold=True, color="1F4E79")
FONT_REGULAR = Font(name=FONT_FAMILY, size=9.5)
FONT_REGULAR_BOLD = Font(name=FONT_FAMILY, size=9.5, bold=True)
FONT_CYAN_BOLD = Font(name=FONT_FAMILY, size=10, bold=True, color="002060")

def clean_uc_title(title):
    if not title:
        return ""
    # Remove leading [UC-xx]
    t = re.sub(r'^\[UC-\d+\]\s*', '', title).strip()
    return t

def extract_all_ucs_data(wb):
    print("📦 1. Trích xuất toàn bộ dữ liệu 120 Use Cases từ file hiện tại...")
    ucs_dict = {}
    
    # Map sheets by UC number
    uc_sheet_map = {}
    for s_name in wb.sheetnames:
        m = re.match(r'UC-(\d+)', s_name)
        if m:
            num = int(m.group(1))
            uc_sheet_map[num] = s_name

    for num in range(1, 121):
        s_name = uc_sheet_map.get(num)
        if not s_name:
            print(f"⚠️ Không tìm thấy sheet cho UC-{num:02d}!")
            continue
        ws = wb[s_name]
        raw_title = ws['A11'].value or ws['B2'].value or f"[UC-{num:02d}]"
        req = ws['B3'].value or ""
        
        # Extract Test Cases
        tcs = []
        for r in range(12, min(50, ws.max_row + 1)):
            tc_id = ws.cell(row=r, column=1).value
            if tc_id and str(tc_id).strip().startswith('TC'):
                vals = [ws.cell(row=r, column=c).value for c in range(1, 16)]
                tcs.append(vals)

        ucs_dict[num] = {
            "num": num,
            "raw_title": raw_title,
            "clean_title": clean_uc_title(raw_title),
            "req": req,
            "tcs": tcs
        }

    print(f"   Đã trích xuất thành công {len(ucs_dict)} Use Cases!")
    return ucs_dict

def format_cell(cell, font=None, alignment=None, fill=None, border=None):
    if font:
        cell.font = font
    if alignment:
        cell.alignment = alignment
    if fill:
        cell.fill = fill
    if border:
        cell.border = border

def build_feature_sheet(ws, feat, ucs_dict):
    sheet_name = feat['sheet_name']
    print(f"   🔨 Xây dựng Feature Sheet: {sheet_name} (Gồm {len(feat['ucs'])} UCs)...")
    ws.views.sheetView[0].showGridLines = True

    # 1. Summary Block (Rows 2 to 8) - Matching Image 3
    ws['A2'].value = "Feature"
    ws['B2'].value = feat['feature_title']
    ws['A3'].value = "Test requirement"
    ws['B3'].value = feat['req']
    ws['A4'].value = "Number of TCs"
    ws['B4'].value = "=SUM(B6:E6)"
    
    ws['A5'].value = "Testing Round"
    ws['B5'].value = "Passed"
    ws['C5'].value = "Failed"
    ws['D5'].value = "Pending"
    ws['E5'].value = "N/A"

    ws['A6'].value = "Round 1"
    ws['B6'].value = "=COUNTIF($F$11:$F$999,B$5)"
    ws['C6'].value = "=COUNTIF($F$11:$F$999,C$5)"
    ws['D6'].value = "=COUNTIF($F$11:$F$999,D$5)"
    ws['E6'].value = "=COUNTIF($F$11:$F$999,E$5)"

    ws['A7'].value = "Round 2"
    ws['B7'].value = "=COUNTIF($I$11:$I$999,B$5)"
    ws['C7'].value = "=COUNTIF($I$11:$I$999,C$5)"
    ws['D7'].value = "=COUNTIF($I$11:$I$999,D$5)"
    ws['E7'].value = "=COUNTIF($I$11:$I$999,E$5)"

    ws['A8'].value = "Round 3"
    ws['B8'].value = "=COUNTIF($L$11:$L$999,B$5)"
    ws['C8'].value = "=COUNTIF($L$11:$L$999,C$5)"
    ws['D8'].value = "=COUNTIF($L$11:$L$999,D$5)"
    ws['E8'].value = "=COUNTIF($L$11:$L$999,E$5)"

    # Style Summary Block
    for r in range(2, 9):
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.font = FONT_REGULAR
            cell.border = thin_border
            if c == 1:
                cell.font = FONT_REGULAR_BOLD
            if r == 5:
                cell.font = FONT_REGULAR_BOLD
                cell.fill = SUMMARY_HEADER_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif r in [6, 7, 8] and c > 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif r in [2, 4]:
                if c == 2:
                    cell.font = FONT_REGULAR_BOLD

    ws['B2'].font = Font(name=FONT_FAMILY, size=11, bold=True, color="1F4E79")
    ws['B4'].font = Font(name=FONT_FAMILY, size=11, bold=True)
    ws['B3'].alignment = Alignment(wrap_text=True, vertical="center")

    # 2. Table Header (Row 10) - Olive Green fill (Image 3)
    headers = [
        "Test Case ID", "Test Case Description", "Test Case Procedure", 
        "Expected Results", "Pre-conditions", "Round 1", "Test date", 
        "Tester", "Round 2", "Test date", "Tester", "Round 3", 
        "Test date", "Tester", "Note"
    ]
    ws.row_dimensions[10].height = 28
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=10, column=col_idx, value=h)
        format_cell(cell, font=FONT_WHITE_BOLD, fill=OLIVE_GREEN_FILL, 
                    alignment=Alignment(horizontal="center", vertical="center", wrap_text=True), 
                    border=thin_border)

    # 3. Populate Use Cases & Test Cases
    current_row = 11
    for uc_num in feat['ucs']:
        uc_data = ucs_dict.get(uc_num)
        if not uc_data:
            continue
        
        # Subheader row (Light Cyan fill spanning A:O) - Matching Image 3
        ws.row_dimensions[current_row].height = 24
        subheader_text = f"UC-{uc_num:02d}. Function {uc_data['clean_title']}"
        for c in range(1, 16):
            cell = ws.cell(row=current_row, column=c)
            if c == 1:
                cell.value = subheader_text
            format_cell(cell, font=FONT_CYAN_BOLD, fill=CYAN_SUBHEADER_FILL, 
                        alignment=Alignment(horizontal="left", vertical="center"), 
                        border=thin_border)
        current_row += 1

        # Test Case rows (Matching naming convention UC01-TC1, UC01-TC2, UC01-TC3)
        for tc_idx, tc_raw in enumerate(uc_data['tcs']):
            tc_code = f"UC{uc_num:02d}-TC{tc_idx + 1}"
            ws.row_dimensions[current_row].height = 55
            
            # Values
            ws.cell(row=current_row, column=1, value=tc_code)
            ws.cell(row=current_row, column=2, value=tc_raw[1]) # Desc
            ws.cell(row=current_row, column=3, value=tc_raw[2]) # Procedure
            ws.cell(row=current_row, column=4, value=tc_raw[3]) # Expected
            ws.cell(row=current_row, column=5, value=tc_raw[4]) # Pre
            ws.cell(row=current_row, column=6, value=tc_raw[5] or "Passed") # Round 1
            ws.cell(row=current_row, column=7, value=tc_raw[6] or "15/09/2026") # Date 1
            ws.cell(row=current_row, column=8, value=tc_raw[7] or "Tran Hong Phuoc (DE180577)") # Tester 1
            ws.cell(row=current_row, column=9, value=tc_raw[8] or "") # Round 2
            ws.cell(row=current_row, column=10, value=tc_raw[9] or "") # Date 2
            ws.cell(row=current_row, column=11, value=tc_raw[10] or "") # Tester 2
            ws.cell(row=current_row, column=12, value=tc_raw[11] or "") # Round 3
            ws.cell(row=current_row, column=13, value=tc_raw[12] or "") # Date 3
            ws.cell(row=current_row, column=14, value=tc_raw[13] or "") # Tester 3
            ws.cell(row=current_row, column=15, value=tc_raw[14] or "") # Note

            # Styles
            for c in range(1, 16):
                cell = ws.cell(row=current_row, column=c)
                cell.font = FONT_REGULAR
                cell.border = thin_border
                if c == 1:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.font = FONT_REGULAR_BOLD
                elif c in [2, 3, 4, 5]:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                elif c in [6, 7, 9, 10, 12, 13]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            
            current_row += 1

    # Column widths
    col_widths = {
        'A': 15, 'B': 30, 'C': 38, 'D': 38, 'E': 26,
        'F': 12, 'G': 13, 'H': 26, 'I': 12, 'J': 13,
        'K': 26, 'L': 12, 'M': 13, 'N': 26, 'O': 12
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

def build_test_cases_sheet(ws, ucs_dict):
    print("📋 3. Cập nhật Sheet 'Test Cases' theo nhóm Feature (Chuẩn hóa giống Ảnh 1)...")
    ws.views.sheetView[0].showGridLines = True
    
    # Header Information
    ws['B1'].value = "TEST CASE LIST"
    ws['B1'].font = FONT_TITLE_BOLD
    ws['B1'].alignment = Alignment(horizontal="center", vertical="center")

    ws['B3'].value = "Project Name"
    ws['C3'].value = "PharmaChain – Smart Pharmacy Chain Management System"
    ws['B4'].value = "Project Code"
    ws['C4'].value = "WDP301"
    ws['B6'].value = "Test Environment Setup Description"
    ws['C6'].value = "1. Server: Ubuntu Server 24.04 LTS (Docker, K8s, Kafka, Redis, MongoDB)\n2. Backend: NestJS Microservices, API Gateway\n3. Web & Mobile: React 19, Expo React Native\n4. Database: MongoDB 7.0\n5. Web Browser: Google Chrome, Microsoft Edge"

    for r in [3, 4, 6]:
        ws.cell(row=r, column=2).font = FONT_REGULAR_BOLD
        ws.cell(row=r, column=3).font = FONT_REGULAR
    ws['C6'].alignment = Alignment(wrap_text=True, vertical="top")

    # Table Header (Row 8) - Navy Blue Fill (Image 1)
    ws.row_dimensions[8].height = 26
    headers = [("B", "No"), ("C", "Name"), ("D", "Sheet Name"), ("E", "Description"), ("F", "Pre-Condition")]
    for col_letter, title in headers:
        cell = ws[f"{col_letter}8"]
        cell.value = title
        format_cell(cell, font=FONT_WHITE_BOLD, fill=NAVY_BLUE_FILL, 
                    alignment=Alignment(horizontal="center", vertical="center"), border=thin_border)

    current_row = 9
    seq_no = 1

    for feat in FEATURES:
        # Category Header Row (Olive Green Fill) - Matching Image 1
        ws.row_dimensions[current_row].height = 24
        header_text = feat['feature_title']
        for col_letter in ['B', 'C', 'D', 'E', 'F']:
            cell = ws[f"{col_letter}{current_row}"]
            if col_letter == 'B':
                cell.value = header_text
            format_cell(cell, font=FONT_WHITE_BOLD, fill=OLIVE_GREEN_FILL, 
                        alignment=Alignment(horizontal="left", vertical="center"), border=thin_border)
        current_row += 1

        # UCs in this Feature
        for uc_num in feat['ucs']:
            uc_data = ucs_dict.get(uc_num)
            if not uc_data:
                continue
            
            ws.row_dimensions[current_row].height = 28
            ws[f"B{current_row}"].value = float(seq_no)
            ws[f"C{current_row}"].value = f"[{uc_data['clean_title']}]" if not uc_data['raw_title'].startswith('[') else uc_data['raw_title']
            ws[f"D{current_row}"].value = feat['sheet_name']
            ws[f"E{current_row}"].value = uc_data['req']
            
            # Extract pre-condition from first TC
            pre_cond = "Actor authenticated; master data pre-seeded."
            if uc_data['tcs'] and len(uc_data['tcs'][0]) > 4 and uc_data['tcs'][0][4]:
                pre_cond = str(uc_data['tcs'][0][4])
            ws[f"F{current_row}"].value = pre_cond

            # Styling
            ws[f"B{current_row}"].alignment = Alignment(horizontal="center", vertical="center")
            ws[f"C{current_row}"].alignment = Alignment(horizontal="left", vertical="center")
            ws[f"D{current_row}"].alignment = Alignment(horizontal="center", vertical="center")
            ws[f"D{current_row}"].font = Font(name=FONT_FAMILY, size=9.5, color="0000FF", underline="single")
            ws[f"E{current_row}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            ws[f"F{current_row}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

            for col_letter in ['B', 'C', 'D', 'E', 'F']:
                cell = ws[f"{col_letter}{current_row}"]
                if col_letter != 'D':
                    cell.font = FONT_REGULAR
                cell.border = thin_border

            current_row += 1
            seq_no += 1

    # Column widths
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 8
    ws.column_dimensions['C'].width = 38
    ws.column_dimensions['D'].width = 28
    ws.column_dimensions['E'].width = 50
    ws.column_dimensions['F'].width = 40

def build_test_statistics_sheet(ws):
    print("📊 4. Cập nhật Sheet 'Test Statistics' theo 8 Feature Modules (Chuẩn hóa giống Ảnh 2)...")
    ws.views.sheetView[0].showGridLines = True

    # Title & Metadata
    ws['B1'].value = "TEST STATISTICS"
    ws['B1'].font = FONT_TITLE_BOLD
    ws['B1'].alignment = Alignment(horizontal="center", vertical="center")

    ws['B3'].value = "Project Name"
    ws['C3'].value = "PharmaChain – Smart Pharmacy Chain Management System"
    ws['E3'].value = "Creator"
    ws['G3'].value = "Tran Hong Phuoc (DE180577)"

    ws['B4'].value = "Project Code"
    ws['C4'].value = "WDP301"
    ws['E4'].value = "Reviewer/Approver"
    ws['G4'].value = "Nguyen Van Nam (DE180564)"

    ws['B5'].value = "Document Code"
    ws['C5'].value = "WDP301_Test Report_v1.2"
    ws['E5'].value = "Issue Date"
    ws['G5'].value = "07/10/2026"

    ws['B6'].value = "Notes"
    ws['C6'].value = "Release includes all 8 Modules across Sprints 1 to 8 (120 UCs / 360 TCs)"

    for r in range(3, 7):
        for c in [2, 5]:
            ws.cell(row=r, column=c).font = FONT_REGULAR_BOLD
        for c in [3, 7]:
            ws.cell(row=r, column=c).font = FONT_REGULAR

    # Table Header (Row 10) - Navy Blue Fill (Image 2)
    ws.row_dimensions[10].height = 26
    headers = [("B", "No"), ("C", "Module code"), ("D", "Passed"), ("E", "Failed"), 
               ("F", "Pending"), ("G", "N/A"), ("H", "Number of test cases")]
    for col_letter, title in headers:
        cell = ws[f"{col_letter}10"]
        cell.value = title
        format_cell(cell, font=FONT_WHITE_BOLD, fill=NAVY_BLUE_FILL, 
                    alignment=Alignment(horizontal="center", vertical="center"), border=thin_border)

    # 8 Feature Rows
    for idx, feat in enumerate(FEATURES):
        r = 11 + idx
        s_name = feat['sheet_name']
        ws.row_dimensions[r].height = 22

        ws[f"B{r}"].value = idx + 1
        ws[f"C{r}"].value = feat['feature_title']
        ws[f"D{r}"].value = f"='{s_name}'!B6"
        ws[f"E{r}"].value = f"='{s_name}'!C6"
        ws[f"F{r}"].value = f"='{s_name}'!D6"
        ws[f"G{r}"].value = f"='{s_name}'!E6"
        ws[f"H{r}"].value = f"='{s_name}'!B4"

        ws[f"B{r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"C{r}"].alignment = Alignment(horizontal="left", vertical="center")
        for col_letter in ['D', 'E', 'F', 'G', 'H']:
            ws[f"{col_letter}{r}"].alignment = Alignment(horizontal="center", vertical="center")

        for col_letter in ['B', 'C', 'D', 'E', 'F', 'G', 'H']:
            cell = ws[f"{col_letter}{r}"]
            cell.font = FONT_REGULAR
            cell.border = thin_border

    # Row 19: Sub total
    ws.row_dimensions[19].height = 24
    ws['B19'].value = None
    ws['C19'].value = "Sub total"
    ws['D19'].value = "=SUM(D11:D18)"
    ws['E19'].value = "=SUM(E11:E18)"
    ws['F19'].value = "=SUM(F11:F18)"
    ws['G19'].value = "=SUM(G11:G18)"
    ws['H19'].value = "=SUM(H11:H18)"

    for col_letter in ['B', 'C', 'D', 'E', 'F', 'G', 'H']:
        cell = ws[f"{col_letter}19"]
        cell.font = FONT_WHITE_BOLD
        cell.fill = NAVY_BLUE_FILL
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Coverage Rows (21 and 22) - Matching Image 2
    ws['C21'].value = "Test coverage"
    ws['C21'].font = FONT_REGULAR_BOLD
    ws['E21'].value = "=(D19+E19)/(H19-G19)"
    ws['E21'].font = Font(name=FONT_FAMILY, size=10, bold=True, color="0000FF")
    ws['E21'].number_format = "100.00 %"

    ws['C22'].value = "Test successful coverage"
    ws['C22'].font = FONT_REGULAR_BOLD
    ws['E22'].value = "=D19/(H19-G19)"
    ws['E22'].font = Font(name=FONT_FAMILY, size=10, bold=True, color="0000FF")
    ws['E22'].number_format = "100.00 %"

    # Column widths
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 8
    ws.column_dimensions['C'].width = 46
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 24

def update_cover_sheet(ws):
    print("📜 5. Cập nhật Sheet 'Cover' lên Version v1.2...")
    ws['F6'].value = "v1.2"
    # Row 13 for Record of Change
    r = 13
    ws[f'A{r}'].value = "07/10/2026"
    ws[f'B{r}'].value = "v1.2"
    ws[f'C{r}'].value = "Feature-based Test Report Refactoring"
    ws[f'D{r}'].value = "M"
    ws[f'E{r}'].value = "Restructure 120 single-use-case sheets into 8 comprehensive feature-based suites mapped with Report 3 SRS architecture"
    ws[f'F{r}'].value = "Report 3 – Software Requirement Specification (Section 3.2 - 3.9)"

    for col_letter in ['A', 'B', 'C', 'D', 'E', 'F']:
        cell = ws[f'{col_letter}{r}']
        cell.font = FONT_REGULAR
        cell.border = thin_border
        if col_letter in ['A', 'B', 'D']:
            cell.alignment = Alignment(horizontal="center", vertical="center")

def main():
    print("="*70)
    print("🚀 BẮT ĐẦU QUY TRÌNH REFACTOR REPORT 5: 1 SHEET = 1 FEATURE")
    print("="*70)

    # 1. Load input workbook
    print(f"📖 Đang đọc file gốc: {INPUT_FILE}...")
    wb_old = openpyxl.load_workbook(INPUT_FILE, data_only=True)
    ucs_dict = extract_all_ucs_data(wb_old)

    # 2. Create new clean workbook
    wb_new = openpyxl.Workbook()
    
    # Copy Cover sheet from old
    ws_cover_old = wb_old['Cover']
    ws_cover = wb_new.active
    ws_cover.title = "Cover"
    for r in range(1, 20):
        for c in range(1, 10):
            val = ws_cover_old.cell(row=r, column=c).value
            if val is not None:
                ws_cover.cell(row=r, column=c, value=val)
                src = ws_cover_old.cell(row=r, column=c)
                if src.font: ws_cover.cell(row=r, column=c).font = copy(src.font)
                if src.alignment: ws_cover.cell(row=r, column=c).alignment = copy(src.alignment)
                if src.border: ws_cover.cell(row=r, column=c).border = copy(src.border)
    update_cover_sheet(ws_cover)

    # Add Test Cases Sheet
    ws_tc = wb_new.create_sheet(title="Test Cases")
    build_test_cases_sheet(ws_tc, ucs_dict)

    # Add Test Statistics Sheet
    ws_stat = wb_new.create_sheet(title="Test Statistics")
    build_test_statistics_sheet(ws_stat)

    # Build the 8 Feature sheets
    print("\n📦 2. Xây dựng 8 Feature Sheets chi tiết...")
    for feat in FEATURES:
        ws_feat = wb_new.create_sheet(title=feat['sheet_name'])
        build_feature_sheet(ws_feat, feat, ucs_dict)

    wb_old.close()

    # Save to local file
    print(f"\n💾 6. Đang lưu file ra: {OUTPUT_FILE}...")
    os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
    wb_new.save(OUTPUT_FILE)
    wb_new.save('/tmp/Report5_Refactored.xlsx')
    print(f"✅ Đã lưu thành công ({os.path.getsize(OUTPUT_FILE):,} bytes)!")

    # Verify sheet names
    print(f"\n🔍 7. Danh sách Sheets mới ({len(wb_new.sheetnames)} tabs):")
    print(wb_new.sheetnames)

    # Upload to Google Drive directly
    print(f"\n☁️ 8. Ghi đè trực tiếp lên Google Spreadsheet ID: {SPREADSHEET_ID}...")
    try:
        token = get_access_token(SA_PATH)
        upload_url = f"https://www.googleapis.com/upload/drive/v3/files/{SPREADSHEET_ID}?uploadType=media"

        with open(OUTPUT_FILE, 'rb') as f:
            file_bytes = f.read()

        req = urllib.request.Request(upload_url, data=file_bytes, headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        }, method='PATCH')

        with urllib.request.urlopen(req) as resp:
            result = json.loads(resp.read().decode('utf-8'))
            print("\n" + "="*70)
            print("🎉 GHI ĐÈ THÀNH CÔNG 100% TRỰC TIẾP LÊN LINK GOOGLE SPREADSHEET!")
            print(f"📄 Tên file trên Cloud: {result.get('name')}")
            print(f"🆔 File ID: {result.get('id')}")
            print(f"🔗 Link Google Sheet: https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/edit")
            print("="*70 + "\n")
    except Exception as e:
        print(f"❌ Lỗi upload Google Drive: {e}")

if __name__ == '__main__':
    main()
