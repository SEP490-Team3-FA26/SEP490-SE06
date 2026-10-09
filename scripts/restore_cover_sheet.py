# scripts/restore_cover_sheet.py
import os
import sys
import json
import urllib.request
import openpyxl
from copy import copy

sys.path.append(os.path.abspath('.'))
sys.path.append(os.path.abspath('.agents/skills/google-docx/scripts'))
from direct_gdrive_sync import get_access_token

SA_PATH = '/Users/tranhongphuoc/.config/gcloud/legacy_credentials/sep490@stone-climate-507417-k4.iam.gserviceaccount.com/adc.json'
SPREADSHEET_ID = '1Fffbgb6gvou2O56s19X4cNvS-o-NOThG'
ORIG_FILE = '/tmp/report5.xlsx'
TARGET_FILE = 'docs/Report5_Test_Report.xlsx'

def restore_cover():
    print("="*60)
    print("🔄 BẮT ĐẦU PHỤC HỒI SHEET COVER VỀ NGUYÊN BẢN CŨ 100%")
    print("="*60)

    print(f"📖 1. Đang đọc file gốc {ORIG_FILE}...")
    wb_orig = openpyxl.load_workbook(ORIG_FILE)
    ws_orig = wb_orig['Cover']

    print(f"📖 2. Đang đọc file hiện tại {TARGET_FILE}...")
    wb_target = openpyxl.load_workbook(TARGET_FILE)
    
    # Remove existing Cover sheet and recreate it at index 0
    if 'Cover' in wb_target.sheetnames:
        del wb_target['Cover']
    
    ws_new = wb_target.create_sheet(title='Cover', index=0)
    ws_new.views.sheetView[0].showGridLines = True

    # 1. Copy column dimensions
    print("📐 3. Sao chép kích thước cột chuẩn...")
    for col_letter, dim in ws_orig.column_dimensions.items():
        if dim.width:
            ws_new.column_dimensions[col_letter].width = dim.width

    # 2. Copy row dimensions
    print("📐 4. Sao chép chiều cao hàng chuẩn...")
    for r in range(1, 15):
        if r in ws_orig.row_dimensions and ws_orig.row_dimensions[r].height:
            ws_new.row_dimensions[r].height = ws_orig.row_dimensions[r].height

    # 3. Copy cells from A1 to G12 (exact original content & styles)
    print("📋 5. Sao chép toàn bộ ô dữ liệu và định dạng gốc (Row 1 -> 12)...")
    for r in range(1, 13):
        for c in range(1, 8):
            src_cell = ws_orig.cell(row=r, column=c)
            dst_cell = ws_new.cell(row=r, column=c)
            
            dst_cell.value = src_cell.value
            if src_cell.has_style:
                if src_cell.font: dst_cell.font = copy(src_cell.font)
                if src_cell.alignment: dst_cell.alignment = copy(src_cell.alignment)
                if src_cell.border: dst_cell.border = copy(src_cell.border)
                if src_cell.fill: dst_cell.fill = copy(src_cell.fill)
                if src_cell.number_format: dst_cell.number_format = src_cell.number_format

    # 4. Copy merged cell ranges
    print("🔗 6. Phục hồi các vùng Merged Cells gốc...")
    for merged_range in ws_orig.merged_cells.ranges:
        print(f"   -> Merge: {merged_range}")
        ws_new.merge_cells(str(merged_range))

    wb_orig.close()

    # 5. Save workbook
    print(f"\n💾 7. Đang lưu file ra: {TARGET_FILE}...")
    wb_target.save(TARGET_FILE)
    wb_target.close()
    print("✅ Đã lưu file thành công!")

    # 6. Upload directly to Google Drive
    print(f"\n☁️ 8. Ghi đè trực tiếp lên Google Spreadsheet ID: {SPREADSHEET_ID}...")
    token = get_access_token(SA_PATH)
    upload_url = f"https://www.googleapis.com/upload/drive/v3/files/{SPREADSHEET_ID}?uploadType=media"

    with open(TARGET_FILE, 'rb') as f:
        file_bytes = f.read()

    req = urllib.request.Request(upload_url, data=file_bytes, headers={
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    }, method='PATCH')

    with urllib.request.urlopen(req) as resp:
        result = json.loads(resp.read().decode('utf-8'))
        print("\n" + "="*60)
        print("🎉 PHỤC HỒI SHEET COVER THÀNH CÔNG 100% TRÊN GOOGLE SPREADSHEET!")
        print(f"📄 Tên file: {result.get('name')}")
        print(f"🆔 File ID: {result.get('id')}")
        print(f"🔗 Link Google Sheet: https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/edit")
        print("="*60 + "\n")

if __name__ == '__main__':
    restore_cover()
