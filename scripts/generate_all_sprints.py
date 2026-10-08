# scripts/generate_all_sprints.py
import os
import sys
import time
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from playwright.sync_api import sync_playwright

sys.path.append(os.path.abspath('.'))
sys.path.append(os.path.abspath('.agents/skills/google-docx/scripts'))

from scripts.generate_sprint_diagrams import generate_sequence_drawio, generate_class_drawio
from direct_gdrive_sync import sync_to_google_drive

SA_PATH = '/Users/tranhongphuoc/.config/gcloud/legacy_credentials/sep490@stone-climate-507417-k4.iam.gserviceaccount.com/adc.json'
GDRIVE_FILE_ID = '1Si1TXNZu-rr6K893iMQ2VcmT4_EucQLK'
DOCX_PATH = 'docs/Report4_Software_Design_Document.docx'

def sanitize_slug(name):
    clean = name.replace(' ', '_').replace('&', 'and').replace('/', '_').replace('-', '_')
    while '__' in clean:
        clean = clean.replace('__', '_')
    return clean

def main():
    print("="*70)
    print("🚀 PHARMA ERP WDP301 - SPRINT DIAGRAM GENERATOR & CLOUD SYNC")
    print("="*70)

    # Active sprints configured for diagram generation (Past Sprints 5-8 archived)
    all_sprints = []

    total_ucs = sum(len(s["ucs"]) for s in all_sprints)
    total_diagrams = total_ucs * 2
    print(f"📦 Total Active Use Cases: {total_ucs}")
    print(f"📊 Total Diagrams to Generate: {total_diagrams} ({total_ucs} Sequence + {total_ucs} Class)")

    os.makedirs('docs/sequence_diagrams', exist_ok=True)
    os.makedirs('docs/class_diagrams', exist_ok=True)

    # 1. Generate XML files
    print("\n📝 1. Generating Draw.io XML files...")
    diagram_tasks = []

    for sprint in all_sprints:
        for uc in sprint["ucs"]:
            slug = f"{uc['id']}_{sanitize_slug(uc['title'])}"
            
            # Sequence
            seq_xml = generate_sequence_drawio(uc)
            seq_drawio = f"docs/sequence_diagrams/{slug}.drawio"
            seq_png = f"docs/sequence_diagrams/{slug}.drawio.png"
            with open(seq_drawio, 'w', encoding='utf-8') as f:
                f.write(seq_xml)
            diagram_tasks.append({
                "type": "seq", "uc": uc, "drawio": seq_drawio, "png": seq_png, "xml": seq_xml
            })

            # Class
            cd_xml = generate_class_drawio(uc)
            cd_drawio = f"docs/class_diagrams/{slug}_ClassDiagram.drawio"
            cd_png = f"docs/class_diagrams/{slug}_ClassDiagram.drawio.png"
            with open(cd_drawio, 'w', encoding='utf-8') as f:
                f.write(cd_xml)
            diagram_tasks.append({
                "type": "class", "uc": uc, "drawio": cd_drawio, "png": cd_png, "xml": cd_xml
            })

    print(f"✅ Generated {len(diagram_tasks)} .drawio XML files!")

    # 2. Render all diagrams with Playwright native export (no clipping, scale=2, border=20)
    print("\n🎨 2. Rendering high-resolution PNGs via Diagrams.net Native Export (Zero-clipping)...")
    render_start = time.time()
    import base64

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width': 2200, 'height': 1600})
        page.goto('file:///tmp/test_single_page.html')
        page.wait_for_function('window.isReady === true', timeout=15000)

        for idx, task in enumerate(diagram_tasks):
            t0 = time.time()
            page.evaluate('window.loadDiagram', task["xml"])
            page.wait_for_timeout(350)
            
            data_url = page.evaluate('window.exportDiagram()')
            if not data_url:
                page.wait_for_timeout(500)
                data_url = page.evaluate('window.exportDiagram()')
                
            if data_url and ',' in data_url:
                b64 = data_url.split(',', 1)[1]
                with open(task["png"], 'wb') as f:
                    f.write(base64.b64decode(b64))
                elapsed = time.time() - t0
                print(f"   [{idx+1:03d}/{len(diagram_tasks)}] {task['uc']['id']} ({task['type'].upper()}): {task['png']} ({elapsed:.2f}s)")
            else:
                print(f"   ❌ FAILED to export {task['uc']['id']} ({task['type'].upper()})")

        browser.close()

    print(f"🎉 Rendered {len(diagram_tasks)} PNGs cleanly in {time.time() - render_start:.1f}s!")

    # 3. Clean and re-append to docx
    print("\n📄 3. Appending diagrams & sections to Report4_Software_Design_Document.docx...")
    doc = docx.Document(DOCX_PATH)
    
    # Truncate paragraphs starting at index 494 so we clean up any old appended sections
    if len(doc.paragraphs) > 494:
        print(f"🧹 Truncating previously appended paragraphs ({len(doc.paragraphs)} -> 494)...")
        for p in doc.paragraphs[494:]:
            p._element.getparent().remove(p._element)

    fig_count = 130
    print(f"ℹ️ Starting figure numbering from Figure {fig_count + 1}...")

    for sprint in all_sprints:
        sec_num = sprint["sec"]
        sprint_title = f"{sec_num} Sprint {sprint['num']}: {sprint['title']}"
        print(f"\n   Adding Section: {sprint_title}")
        doc.add_paragraph(sprint_title, style='Heading 3')

        sub_idx = 1
        for uc in sprint["ucs"]:
            slug = f"{uc['id']}_{sanitize_slug(uc['title'])}"
            cd_png = f"docs/class_diagrams/{slug}_ClassDiagram.drawio.png"
            seq_png = f"docs/sequence_diagrams/{slug}.drawio.png"

            # 3.X.Y Class Diagram
            cd_heading = f"{sec_num}.{sub_idx} Class Diagram for {uc['title']} ({uc['id']})"
            doc.add_paragraph(cd_heading, style='Heading 4')
            
            p_img_cd = doc.add_paragraph()
            p_img_cd.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_cd = p_img_cd.add_run()
            r_cd.add_picture(cd_png, width=Inches(6.2))

            fig_count += 1
            p_cap_cd = doc.add_paragraph()
            p_cap_cd.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rc_cd = p_cap_cd.add_run(f"Figure {fig_count}. Class diagram for {uc['title_short']}")
            rc_cd.bold = True
            rc_cd.font.name = "Arial"
            rc_cd.font.size = Pt(9.5)

            sub_idx += 1

            # 3.X.Y+1 Sequence Diagram
            seq_heading = f"{sec_num}.{sub_idx} Sequence Diagram for {uc['title']} ({uc['id']})"
            doc.add_paragraph(seq_heading, style='Heading 4')
            
            p_img_seq = doc.add_paragraph()
            p_img_seq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_seq = p_img_seq.add_run()
            r_seq.add_picture(seq_png, width=Inches(6.2))

            fig_count += 1
            p_cap_seq = doc.add_paragraph()
            p_cap_seq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rc_seq = p_cap_seq.add_run(f"Figure {fig_count}. Sequence diagram for {uc['title_short']}")
            rc_seq.bold = True
            rc_seq.font.name = "Arial"
            rc_seq.font.size = Pt(9.5)

            sub_idx += 1

    doc.save(DOCX_PATH)
    print(f"\n💾 Saved updated document to {DOCX_PATH} with total {fig_count} figures!")

    # 4. Overwrite Google Doc via Direct Google Drive Sync
    print("\n☁️ 4. Syncing directly to Google Docs Cloud Document...")
    sync_to_google_drive(GDRIVE_FILE_ID, DOCX_PATH, SA_PATH)
    print("\n🎉 ALL TASKS COMPLETED SUCCESSFULLY!")

if __name__ == '__main__':
    main()
