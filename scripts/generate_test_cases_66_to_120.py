# scripts/generate_test_cases_66_to_120.py
import os
import sys
import json
import time
import openpyxl
from copy import copy
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import urllib.request

sys.path.append(os.path.abspath('.'))
sys.path.append(os.path.abspath('.agents/skills/google-docx/scripts'))
from direct_gdrive_sync import get_access_token

SA_PATH = '/Users/tranhongphuoc/.config/gcloud/legacy_credentials/sep490@stone-climate-507417-k4.iam.gserviceaccount.com/adc.json'
SPREADSHEET_ID = '1Fffbgb6gvou2O56s19X4cNvS-o-NOThG'
LOCAL_INPUT = '/tmp/Report5_Test_Report.xlsx'
LOCAL_OUTPUT = 'docs/Report5_Test_Report.xlsx'

UCS_66_TO_120 = [
    # SPRINT 5
    {
        "id": "UC-66", "num": 66,
        "title": "RFM Customer Segmentation & Churn Analytics",
        "sheet": "UC-66 RFM Customer Segment",
        "actor": "Marketing Manager, Admin",
        "req": "System classifies customer base into Recency, Frequency, Monetary (RFM) cohorts and flags churn risk customers for targeted re-engagement campaigns.",
        "pre": "Actor 'Marketing Manager, Admin' authenticated; customer purchase order history pre-seeded."
    },
    {
        "id": "UC-67", "num": 67,
        "title": "Branch Cash Drawer Management",
        "sheet": "UC-67 Branch Cash Drawer",
        "actor": "Pharmacist, Branch Manager",
        "req": "Allows opening, tracking shift cash float, logging petty cash payouts, and verifying cash counts against POS sales.",
        "pre": "Actor 'Pharmacist, Branch Manager' authenticated; POS till terminal active."
    },
    {
        "id": "UC-68", "num": 68,
        "title": "Pharmacist OCR Prescription Verification",
        "sheet": "UC-68 OCR Prescription Verif",
        "actor": "Pharmacist",
        "req": "Pharmacist uploads doctor prescription photo; OCR extracts active ingredients, dosage, and validates against hospital registry.",
        "pre": "Actor 'Pharmacist' authenticated; prescription image file ready for verification."
    },
    {
        "id": "UC-69", "num": 69,
        "title": "Prescription OCR Audit History",
        "sheet": "UC-69 OCR Audit History",
        "actor": "Branch Manager, Admin",
        "req": "Audits past OCR scans, pharmacist manual adjustments, compliance overrides, and doctor registration references.",
        "pre": "Actor 'Branch Manager, Admin' authenticated; historical prescription OCR records exist."
    },
    {
        "id": "UC-70", "num": 70,
        "title": "Pharmacist Consultation Voice Recording History",
        "sheet": "UC-70 Voice Record History",
        "actor": "Branch Manager, Admin",
        "req": "Stores and manages audio recordings of customer drug consultations for QA, dispute settlement, and pharmacist compliance tracking.",
        "pre": "Actor 'Branch Manager, Admin' authenticated; audio consultation recordings stored on AWS S3."
    },
    {
        "id": "UC-71", "num": 71,
        "title": "Branch Cash Flow Reporting",
        "sheet": "UC-71 Branch Cash Flow Report",
        "actor": "Branch Manager, Accountant",
        "req": "Generates daily cash in, cash out, card settlements, VietQR totals, and reconciles against branch bank deposits.",
        "pre": "Actor 'Branch Manager, Accountant' authenticated; closed shift ledger data present."
    },
    {
        "id": "UC-72", "num": 72,
        "title": "Branch & Chain Cash Flow Analysis",
        "sheet": "UC-72 Chain Cash Flow Analys",
        "actor": "Head Manager, Finance Officer",
        "req": "Aggregates cash flow across all regional branches, identifying net burn rate, capital liquidity, and operating margin trends.",
        "pre": "Actor 'Head Manager, Finance Officer' authenticated; multi-branch financial streams synchronized."
    },
    {
        "id": "UC-73", "num": 73,
        "title": "Fixed Cost & Expense Governance",
        "sheet": "UC-73 Fixed Cost Governance",
        "actor": "Head Manager, Accountant",
        "req": "Manages branch fixed overhead (store rental, power/water utilities, monthly staff payroll, equipment depreciation).",
        "pre": "Actor 'Head Manager, Accountant' authenticated; monthly store expense allocations configured."
    },
    # SPRINT 6
    {
        "id": "UC-74", "num": 74,
        "title": "Time-limited Flash Sale Management",
        "sheet": "UC-74 Flash Sale Management",
        "actor": "Marketing Manager, Admin",
        "req": "Configures countdown flash sales with stock quota allocations, maximum cart quantity per user, and automatic price reversion upon expiry.",
        "pre": "Actor 'Marketing Manager, Admin' authenticated; campaign inventory quotas defined."
    },
    {
        "id": "UC-75", "num": 75,
        "title": "Advanced Promotional Voucher & Combo",
        "sheet": "UC-75 Promo Voucher and Combo",
        "actor": "Marketing Manager, Admin",
        "req": "Creates conditional multi-tier vouchers (Buy X Get Y free, cross-category bundle discounts, minimum bill thresholds).",
        "pre": "Actor 'Marketing Manager, Admin' authenticated; product master catalog accessible."
    },
    {
        "id": "UC-76", "num": 76,
        "title": "Automated Marketing Email Campaigns",
        "sheet": "UC-76 Marketing Email Camp",
        "actor": "Marketing Manager, Admin",
        "req": "Sends automated email newsletters, coupon drops, and abandoned cart reminder sequences based on customer RFM segments.",
        "pre": "Actor 'Marketing Manager, Admin' authenticated; SMTP/SendGrid mail service connected."
    },
    {
        "id": "UC-77", "num": 77,
        "title": "Customer Product Review & Rating",
        "sheet": "UC-77 Customer Product Review",
        "actor": "Customer, Pharmacist",
        "req": "Allows verified purchasers to submit 1-5 star ratings, photo reviews, and side-effect feedback for pharmaceutical items.",
        "pre": "Actor 'Customer' authenticated with verified delivered order status."
    },
    {
        "id": "UC-78", "num": 78,
        "title": "Import Medicine Auxiliary Label & Barcode Printing",
        "sheet": "UC-78 Aux Label and Barcode",
        "actor": "Warehouse Staff",
        "req": "Generates and prints Ministry of Health compliant secondary auxiliary Vietnamese labels and thermal EAN-13 barcodes for imported items.",
        "pre": "Actor 'Warehouse Staff' authenticated; thermal label printer connected."
    },
    {
        "id": "UC-79", "num": 79,
        "title": "Multi-Supplier Bulk RFQ Quotation",
        "sheet": "UC-79 Multi-Supplier Bulk RFQ",
        "actor": "Procurement Officer",
        "req": "Dispatches Request for Quotation (RFQ) packages to multiple GDP pharma suppliers simultaneously and parses supplier bid responses.",
        "pre": "Actor 'Procurement Officer' authenticated; supplier master directory pre-seeded."
    },
    {
        "id": "UC-80", "num": 80,
        "title": "AI Symptom Consultation Chatbot",
        "sheet": "UC-80 AI Symptom Chatbot",
        "actor": "Customer",
        "req": "Provides automated interactive AI consultation based on symptoms, recommends OTC relief medicines, and refers to doctors if severe.",
        "pre": "Actor 'Customer' authenticated; Gemini AI engine online."
    },
    {
        "id": "UC-81", "num": 81,
        "title": "AI Voice Medicine Search",
        "sheet": "UC-81 AI Voice Medicine Search",
        "actor": "Customer, Pharmacist",
        "req": "Recognizes spoken medical queries in Vietnamese/English, transliterating brand names to generic INN drug names with autocomplete.",
        "pre": "Microphone permissions granted; WebSpeech/Whisper API operational."
    },
    {
        "id": "UC-82", "num": 82,
        "title": "IoT Central Warehouse Telemetry & Alerts",
        "sheet": "UC-82 IoT Warehouse Telemetry",
        "actor": "Warehouse Staff, Admin",
        "req": "Captures real-time temperature, humidity, and power backup telemetry from GSP warehouse cold storage sensors with breach alarms.",
        "pre": "IoT sensor gateway emitting MQTT telemetry to Kafka topic 'sensor.telemetry.ingest'."
    },
    {
        "id": "UC-83", "num": 83,
        "title": "Marketing Campaign ROI Analytics",
        "sheet": "UC-83 Campaign ROI Analytics",
        "actor": "Marketing Manager, Head Manager",
        "req": "Calculates Customer Acquisition Cost (CAC), Return on Ad Spend (ROAS), and gross revenue uplift attributed to specific promo codes.",
        "pre": "Actor 'Marketing Manager' authenticated; UTM campaign tracking parameters recorded."
    },
    # SPRINT 7
    {
        "id": "UC-84", "num": 84,
        "title": "Shift Handover & Cash Balancing",
        "sheet": "UC-84 Shift Handover Cash",
        "actor": "Pharmacist, Branch Manager",
        "req": "Performs formal shift handover, counting drawer bills, reconciling card/VietQR receipts, and obtaining digital handover signatures.",
        "pre": "Active pharmacist shift nearing conclusion; POS transactions locked for count."
    },
    {
        "id": "UC-85", "num": 85,
        "title": "Order & Transaction Cancellation",
        "sheet": "UC-85 Order Cancellation",
        "actor": "Pharmacist, Branch Manager",
        "req": "Processes transaction voiding, order cancellation with mandatory reason codes, stock replenishment rollbacks, and audit logging.",
        "pre": "Actor 'Pharmacist, Branch Manager' authenticated; target order in pending/unsettled state."
    },
    {
        "id": "UC-86", "num": 86,
        "title": "Distance-based Delivery Fee Calculation",
        "sheet": "UC-86 Delivery Fee Calc",
        "actor": "Customer, Pharmacist",
        "req": "Computes shipping charges using Google Maps Distance Matrix API based on customer coordinate distance to closest fulfilling branch.",
        "pre": "Customer shipping address coordinates validated."
    },
    {
        "id": "UC-87", "num": 87,
        "title": "Barcode Scanning Goods Receipt",
        "sheet": "UC-87 Barcode Goods Receipt",
        "actor": "Warehouse Staff",
        "req": "Scans GS1-128 / DataMatrix barcodes on supplier cartons, auto-populating batch number, manufacturing date, and expiry date into GRN.",
        "pre": "Physical shipment delivered; barcode scanner paired."
    },
    {
        "id": "UC-88", "num": 88,
        "title": "Barcode Scanning Stock Dispatch",
        "sheet": "UC-88 Barcode Stock Dispatch",
        "actor": "Warehouse Staff",
        "req": "Scans picking item barcodes to verify SKU match against outgoing transfer orders or courier delivery shipments before seal.",
        "pre": "Stock dispatch order approved; picking items staged."
    },
    {
        "id": "UC-89", "num": 89,
        "title": "Barcode Scanner Inventory Audit",
        "sheet": "UC-89 Barcode Stock Audit",
        "actor": "Warehouse Staff, Pharmacist",
        "req": "Conducts cyclical GSP shelf stocktaking with hand-held wireless barcode terminal, reporting instant variance against ERP system stock.",
        "pre": "Stocktaking session initiated; warehouse zone assigned."
    },
    {
        "id": "UC-90", "num": 90,
        "title": "GSP Bin & Rack Warehouse Location Management",
        "sheet": "UC-90 GSP Bin and Rack Mgmt",
        "actor": "Warehouse Staff, Admin",
        "req": "Manages 2D/3D warehouse storage map, allocating fast-moving items to ergonomic lower tiers and cold-chain items to monitored chilled bins.",
        "pre": "Warehouse zone layout pre-configured."
    },
    {
        "id": "UC-91", "num": 91,
        "title": "SKU Code Generation & GS1 Barcode Generator",
        "sheet": "UC-91 SKU and GS1 Generator",
        "actor": "Inventory Manager, Admin",
        "req": "Generates compliant GS1 EAN-13 barcodes with checksum digits and standardized internal SKU hierarchy for newly registered medicines.",
        "pre": "Medicine catalog entry created with active ingredient classification."
    },
    {
        "id": "UC-92", "num": 92,
        "title": "Multi-Unit of Measure Conversion",
        "sheet": "UC-92 Multi-UOM Conversion",
        "actor": "Inventory Manager, Pharmacist",
        "req": "Defines conversion coefficients (e.g. 1 Box = 10 Blisters = 100 Tablets) allowing retail dispensing by tablet and purchase by master carton.",
        "pre": "Medicine catalog item active; base packaging unit configured."
    },
    {
        "id": "UC-93", "num": 93,
        "title": "Branch-specific Pricing Policy",
        "sheet": "UC-93 Branch Pricing Policy",
        "actor": "Head Manager, Pricing Admin",
        "req": "Configures regional price overrides, tier markups, and localized promotional margins per branch location.",
        "pre": "Branch network registered; base catalog pricing established."
    },
    {
        "id": "UC-94", "num": 94,
        "title": "Negative & Discrepancy Stock Alert",
        "sheet": "UC-94 Negative Stock Alert",
        "actor": "Inventory Manager, Branch Manager",
        "req": "Intercepts abnormal negative inventory deductions, triggering audit quarantine and dispatching instant alert to store manager.",
        "pre": "Real-time stock ledger service running."
    },
    {
        "id": "UC-95", "num": 95,
        "title": "AI Cross-sell & Bundled Medicine Recommendation",
        "sheet": "UC-95 AI Cross-sell Bundles",
        "actor": "Pharmacist, Customer",
        "req": "AI suggests complementary OTC medicines (e.g. Vitamin C + Paracetamol) based on clinical association rules and cross-selling models.",
        "pre": "Active shopping cart contains at least one primary SKU."
    },
    # SPRINT 8
    {
        "id": "UC-96", "num": 96,
        "title": "Customer Wishlist & Stock Back-In Notification",
        "sheet": "UC-96 Wishlist Stock Alert",
        "actor": "Customer",
        "req": "Allows customers to bookmark out-of-stock specialty drugs and dispatches instant push notification when fresh batches are receipted.",
        "pre": "Actor 'Customer' authenticated; item currently at zero inventory."
    },
    {
        "id": "UC-97", "num": 97,
        "title": "Rare Medicine Pre-Order Booking",
        "sheet": "UC-97 Rare Med Pre-Order",
        "actor": "Customer, Pharmacist",
        "req": "Facilitates deposit pre-orders for rare orphan drugs, initiating automated procurement back-to-back with GDP certified importers.",
        "pre": "Customer prescription validated; item flagged as special import."
    },
    {
        "id": "UC-98", "num": 98,
        "title": "Post-Purchase CSAT Feedback Survey",
        "sheet": "UC-98 CSAT Feedback Survey",
        "actor": "Customer, Customer Care",
        "req": "Dispatches 3-question Net Promoter Score (NPS) and Customer Satisfaction (CSAT) survey following order delivery or counter consultation.",
        "pre": "Order status transitioned to 'Delivered' / 'Completed'."
    },
    {
        "id": "UC-99", "num": 99,
        "title": "Health Quiz Gamification & Point Rewards",
        "sheet": "UC-99 Health Gamification",
        "actor": "Customer",
        "req": "Presents daily health trivia quizzes on mobile app, rewarding verified loyalty points upon correct answers to boost customer engagement.",
        "pre": "Customer mobile profile active; daily quiz pool available."
    },
    {
        "id": "UC-100", "num": 100,
        "title": "Delivery Proof & Shipper QR Scanner",
        "sheet": "UC-100 Delivery Proof QR",
        "actor": "Shipper, Customer",
        "req": "Internal pharmacy shipper scans recipient QR code and captures e-signature photo as cryptographic proof of handover.",
        "pre": "Shipper assigned to dispatch route; mobile camera enabled."
    },
    {
        "id": "UC-101", "num": 101,
        "title": "Social Pixel & Omnichannel Attribution Tracking",
        "sheet": "UC-101 Omnichannel Attribution",
        "actor": "Marketing Manager, Admin",
        "req": "Tracks Facebook CAPI, TikTok Pixel, and Google Ads conversions across web portal, mobile app, and in-store POS checkouts.",
        "pre": "Tracking pixels initialized; privacy consent recorded."
    },
    {
        "id": "UC-102", "num": 102,
        "title": "Narcotic & Restricted Medicine Compliance Log",
        "sheet": "UC-102 Restricted Drug Log",
        "actor": "Pharmacist, Chief Pharmacist",
        "req": "Maintains tamper-proof Ministry of Health special register for habit-forming narcotics, recording patient identity number and prescriber licence.",
        "pre": "Prescription contains Schedule II/III restricted substances."
    },
    {
        "id": "UC-103", "num": 103,
        "title": "Supplier GDP Master Contract & Terms",
        "sheet": "UC-103 Supplier GDP Contract",
        "actor": "Procurement Officer, Legal",
        "req": "Archives Good Distribution Practice (GDP) certificates, credit period terms, defect recall clauses, and annual volume rebates per supplier.",
        "pre": "Supplier profile verified; legal documents uploaded."
    },
    {
        "id": "UC-104", "num": 104,
        "title": "Supplier Payment Voucher Settlement",
        "sheet": "UC-104 Supplier Payment Settle",
        "actor": "Accountant, Finance Officer",
        "req": "Matches 3-way PO, GRN, and vendor invoice before generating electronic bank transfer voucher for supplier accounts payable.",
        "pre": "GRN approved; vendor e-invoice uploaded."
    },
    {
        "id": "UC-105", "num": 105,
        "title": "AI Supply Chain Disruption & Delay Risk Forecast",
        "sheet": "UC-105 AI Supply Chain Risk",
        "actor": "Procurement Officer, Head Manager",
        "req": "Analyzes port congestion, seasonal weather trends, and vendor delivery punctuality to predict stockout vulnerabilities 30 days in advance.",
        "pre": "Historical vendor lead-time dataset loaded."
    },
    {
        "id": "UC-106", "num": 106,
        "title": "Digital Health Passport & Electronic Health Record",
        "sheet": "UC-106 Health Passport EHR",
        "actor": "Customer, Pharmacist",
        "req": "Encrypted digital wallet storing chronic medication schedules, vaccination logs, and allergies accessible across all chain branches.",
        "pre": "Customer biometric authorization granted."
    },
    {
        "id": "UC-107", "num": 107,
        "title": "AI Allergy & Contraindication Cross-check",
        "sheet": "UC-107 AI Allergy Cross-check",
        "actor": "Pharmacist",
        "req": "Cross-references prescribed items against patient chronic history and penicillin/sulfa allergies, warning pharmacists before sale.",
        "pre": "Customer health profile linked; items in checkout cart."
    },
    {
        "id": "UC-108", "num": 108,
        "title": "Prescription Adherence Tracking & Physician Report",
        "sheet": "UC-108 Adherence Tracking",
        "actor": "Customer, Pharmacist",
        "req": "Monitors customer pill intake adherence via mobile reminders and compiles monthly compliance score reports for treating physicians.",
        "pre": "Patient enrolled in chronic disease management program."
    },
    {
        "id": "UC-109", "num": 109,
        "title": "Multilingual AI Voice Commerce & Ordering",
        "sheet": "UC-109 Voice Commerce Order",
        "actor": "Customer",
        "req": "Allows elderly or visually impaired customers to order repeat prescriptions using conversational Vietnamese voice prompts.",
        "pre": "Customer phone verified; microphone access allowed."
    },
    {
        "id": "UC-110", "num": 110,
        "title": "AI Best Procurement Price & Supplier Routing",
        "sheet": "UC-110 AI Best Price Routing",
        "actor": "Procurement Officer",
        "req": "Optimizes purchase order splitting across competing suppliers to achieve minimum total landed cost considering delivery lead time.",
        "pre": "Pending bulk purchase requisition active."
    },
    {
        "id": "UC-111", "num": 111,
        "title": "Pharmacist Sales KPI & Target Tracking",
        "sheet": "UC-111 Pharmacist Sales KPI",
        "actor": "Branch Manager, Pharmacist",
        "req": "Tracks daily counter turnover, OTC attachment rates, customer review scores, and commission bonuses against assigned branch quotas.",
        "pre": "Monthly branch sales targets configured."
    },
    {
        "id": "UC-112", "num": 112,
        "title": "Disaster Recovery Backup & Manual Restore Log",
        "sheet": "UC-112 Disaster Recovery Log",
        "actor": "System Admin",
        "req": "Initiates point-in-time MongoDB Atlas snapshots, verifies offsite S3 cold storage archives, and audits emergency restore trial runs.",
        "pre": "Admin master credentials authenticated; backup service online."
    },
    {
        "id": "UC-113", "num": 113,
        "title": "Biometric Mobile Authentication (FaceID/Fingerprint)",
        "sheet": "UC-113 Biometric Mobile Auth",
        "actor": "Customer, Pharmacist",
        "req": "Integrates iOS FaceID and Android BiometricPrompt for cryptographic token exchange without manual password input.",
        "pre": "Device hardware supports biometric sensors; app enrolled."
    },
    {
        "id": "UC-114", "num": 114,
        "title": "IoT Device Management & Sensor Calibration",
        "sheet": "UC-114 IoT Sensor Calibration",
        "actor": "IoT Technician, Admin",
        "req": "Registers LoRaWAN/Zigbee temperature probes, schedules annual GSP calibration cycles, and sets drift offset coefficients.",
        "pre": "Sensor hardware UUID provisioned on network."
    },
    # EXTENDED GOVERNANCE UC-115 to UC-120
    {
        "id": "UC-115", "num": 115,
        "title": "Offline Data Synchronization & Conflict Resolution",
        "sheet": "UC-115 Offline Data Sync",
        "actor": "Pharmacist, System Admin",
        "req": "Allows POS counter to continue dispensing sales during WAN network outages with local SQLite storage, auto-syncing with Kafka upon reconnection.",
        "pre": "Offline cache database operational on local POS terminal."
    },
    {
        "id": "UC-116", "num": 116,
        "title": "Retail Customer Debt Limit & Aging Schedule",
        "sheet": "UC-116 Customer Debt Limit",
        "actor": "Accountant, Branch Manager",
        "req": "Tracks credit receivables for regular chronic patients, enforcing maximum credit balances and generating 30/60/90-day aging debt notices.",
        "pre": "Customer credit account approved by store manager."
    },
    {
        "id": "UC-117", "num": 117,
        "title": "Real-time Shipper GPS Fleet Tracking",
        "sheet": "UC-117 Shipper GPS Tracking",
        "actor": "Branch Manager, Customer",
        "req": "Displays live telemetry map of internal pharmacy delivery motorbikes, estimating real-time arrival windows for urgent home prescriptions.",
        "pre": "Shipper app transmitting continuous GPS coordinates."
    },
    {
        "id": "UC-118", "num": 118,
        "title": "Recurring Medication Auto-Reorder Management",
        "sheet": "UC-118 Auto-Reorder Mgmt",
        "actor": "Customer, Pharmacist",
        "req": "Schedules monthly repeat prescription dispensations for chronic conditions (hypertension, diabetes), auto-preparing order 3 days in advance.",
        "pre": "Customer enrolled in recurring prescription auto-refill plan."
    },
    {
        "id": "UC-119", "num": 119,
        "title": "Import Medicine Auxiliary Label Generator",
        "sheet": "UC-119 Aux Label Generator",
        "actor": "Inventory Manager, Warehouse Staff",
        "req": "Translates foreign packaging drug monographs into standardized Vietnamese supplementary labels compliant with Circular 01/2018/TT-BYT.",
        "pre": "Foreign drug registration dossier and visa number registered."
    },
    {
        "id": "UC-120", "num": 120,
        "title": "System Multilingual Localization (VI/EN)",
        "sheet": "UC-120 Multilingual Localiz",
        "actor": "Customer, System Admin",
        "req": "Provides seamless real-time UI localization between Vietnamese and English across web portal, mobile app, and receipt printing.",
        "pre": "i18n translation dictionary loaded for target language."
    },
]

def update_test_report():
    print("="*70)
    print("🚀 PHARMA ERP WDP301 - EXTENDING TEST REPORT TO UC-120")
    print("="*70)

    # 1. Load existing workbook
    print(f"📖 1. Loading workbook from {LOCAL_INPUT}...")
    wb = openpyxl.load_workbook(LOCAL_INPUT)

    # 2. Update UC-60 to UC-65: Mark Round 1 as 'Pending' (as specified: pre-generated, not yet tested)
    print("🔄 2. Synchronizing UC-60 to UC-65 (Setting Round 1 to 'Pending')...")
    for uc_num in range(60, 66):
        # Find matching sheet
        for sheet_name in wb.sheetnames:
            if sheet_name.startswith(f"UC-{uc_num:02d}") or sheet_name.startswith(f"UC-{uc_num}"):
                ws = wb[sheet_name]
                for r in [12, 13, 14]:
                    ws[f'F{r}'].value = 'Pending'
                    ws[f'G{r}'].value = None
                    ws[f'H{r}'].value = None
                    ws[f'I{r}'].value = None
                    ws[f'J{r}'].value = None
                    ws[f'K{r}'].value = None
                print(f"   Updated {sheet_name} Round 1 -> 'Pending'")
                break

    # 3. Create new sheets for UC-66 to UC-120
    print(f"\n📑 3. Creating 55 new test case sheets (UC-66 to UC-120)...")
    template_sheet = wb['UC-65 Receive notifications']
    
    for idx, uc in enumerate(UCS_66_TO_120):
        t0 = time.time()
        sheet_title = uc['sheet']
        if sheet_title in wb.sheetnames:
            ws = wb[sheet_title]
        else:
            ws = wb.copy_worksheet(template_sheet)
            ws.title = sheet_title

        # Update metadata
        ws['B2'].value = f"[{uc['id']}] {uc['title']}"
        ws['B3'].value = uc['req']
        ws['A11'].value = f"[{uc['id']}] {uc['title']}"

        # Row 12: Happy Path
        ws['A12'].value = f"TC-{uc['num']}"
        ws['B12'].value = f"Verify successful execution of '{uc['title']}' (Happy Path)."
        ws['C12'].value = f"1. Access system with role '{uc['actor']}'.\n2. Navigate to '{uc['title']}' function.\n3. Enter valid and complete required parameters.\n4. Submit and verify system response."
        ws['D12'].value = f"1. System validates inputs and processes successfully.\n2. Confirmation message displayed.\n3. Data persisted to MongoDB and synchronized via Kafka.\n4. Screen state transitions to corresponding post-action view."
        ws['E12'].value = uc['pre']
        ws['F12'].value = 'Pending'
        ws['G12'].value = None
        ws['H12'].value = None
        ws['I12'].value = None
        ws['J12'].value = None
        ws['K12'].value = None

        # Row 13: Negative Path
        ws['A13'].value = f"TC-{uc['num']}-02"
        ws['B13'].value = f"Verify validation and error handling for '{uc['title']}' with invalid/missing inputs (Negative Path)."
        ws['C13'].value = f"1. Access system with role '{uc['actor']}'.\n2. Navigate to '{uc['title']}' function.\n3. Leave mandatory fields blank or enter invalid data format.\n4. Attempt to submit.\n5. Observe validation error warnings."
        ws['D13'].value = f"1. System blocks submission.\n2. Clear error alerts displayed for invalid fields.\n3. No corrupted data is written to Database.\n4. Form inputs preserved for correction."
        ws['E13'].value = f"Actor '{uc['actor']}' authenticated."
        ws['F13'].value = 'Pending'
        ws['G13'].value = None
        ws['H13'].value = None
        ws['I13'].value = None
        ws['J13'].value = None
        ws['K13'].value = None

        # Row 14: Edge / Security Path
        ws['A14'].value = f"TC-{uc['num']}-03"
        ws['B14'].value = f"Verify boundary values, concurrency, or security constraints for '{uc['title']}' (Edge Path)."
        ws['C14'].value = f"1. Access '{uc['title']}' function.\n2. Test boundary lengths, duplicate submissions, or unauthorized role access.\n3. Observe system defense and audit logging."
        ws['D14'].value = f"1. Boundary conditions handled gracefully.\n2. Concurrency locks prevent duplicate processing.\n3. RBAC security permissions enforced and audit log recorded."
        ws['E14'].value = f"Actor '{uc['actor']}' authenticated."
        ws['F14'].value = 'Pending'
        ws['G14'].value = None
        ws['H14'].value = None
        ws['I14'].value = None
        ws['J14'].value = None
        ws['K14'].value = None

        print(f"   [{idx+1:02d}/55] Created sheet: {sheet_title} ({time.time() - t0:.2f}s)")

    # 4. Update 'Test Cases' master index sheet (gid=1077195425)
    print("\n📋 4. Updating 'Test Cases' master index sheet...")
    tc_sheet = wb['Test Cases']
    
    # Rows 73 to 127
    start_row = 73
    for idx, uc in enumerate(UCS_66_TO_120):
        r = start_row + idx
        tc_sheet[f'B{r}'].value = uc['num']
        tc_sheet[f'C{r}'].value = f"[{uc['id']}] {uc['title']}"
        tc_sheet[f'D{r}'].value = uc['sheet']
        tc_sheet[f'E{r}'].value = uc['req']
        tc_sheet[f'F{r}'].value = uc['pre']

        # Copy style from row 72
        for col_letter in ['B', 'C', 'D', 'E', 'F']:
            src_cell = tc_sheet[f'{col_letter}72']
            dst_cell = tc_sheet[f'{col_letter}{r}']
            if src_cell.font:
                dst_cell.font = copy(src_cell.font)
            if src_cell.alignment:
                dst_cell.alignment = copy(src_cell.alignment)
            if src_cell.border:
                dst_cell.border = copy(src_cell.border)

    print(f"   Added rows 73 to 127 in 'Test Cases' sheet!")

    # 5. Update 'Test Statistics' sheet
    print("\n📊 5. Updating 'Test Statistics' sheet...")
    stat_sheet = wb['Test Statistics']

    # Insert rows from 76 to 130 for UC-66 to UC-120
    for idx, uc in enumerate(UCS_66_TO_120):
        r = 76 + idx
        s_title = uc['sheet']
        stat_sheet[f'B{r}'].value = float(uc['num'])
        stat_sheet[f'C{r}'].value = f"='{s_title}'!B2"
        stat_sheet[f'D{r}'].value = f"='{s_title}'!B6"
        stat_sheet[f'E{r}'].value = f"='{s_title}'!C6"
        stat_sheet[f'F{r}'].value = f"='{s_title}'!D6"
        stat_sheet[f'G{r}'].value = f"='{s_title}'!E6"
        stat_sheet[f'H{r}'].value = f"='{s_title}'!B4"

        # Apply styles from row 75
        for col_letter in ['B', 'C', 'D', 'E', 'F', 'G', 'H']:
            src_cell = stat_sheet[f'{col_letter}75']
            dst_cell = stat_sheet[f'{col_letter}{r}']
            if src_cell.font:
                dst_cell.font = copy(src_cell.font)
            if src_cell.alignment:
                dst_cell.alignment = copy(src_cell.alignment)
            if src_cell.border:
                dst_cell.border = copy(src_cell.border)

    # Sub total row at 131
    stat_sheet['B131'].value = None
    stat_sheet['C131'].value = 'Sub total'
    stat_sheet['D131'].value = '=SUM(D11:D130)'
    stat_sheet['E131'].value = '=SUM(E11:E130)'
    stat_sheet['F131'].value = '=SUM(F11:F130)'
    stat_sheet['G131'].value = '=SUM(G11:G130)'
    stat_sheet['H131'].value = '=SUM(H11:H130)'

    # Formatting for totals
    bold_font = Font(name='Tahoma', size=10, bold=True)
    stat_sheet['C131'].font = bold_font
    for col in ['D', 'E', 'F', 'G', 'H']:
        stat_sheet[f'{col}131'].font = bold_font
        stat_sheet[f'{col}131'].number_format = '#,##0'

    # Coverage rows
    stat_sheet['C133'].value = 'Test coverage'
    stat_sheet['C133'].font = bold_font
    stat_sheet['E133'].value = '=(D131+E131)/(H131-G131)'
    stat_sheet['E133'].font = bold_font
    stat_sheet['E133'].number_format = '0.00%'

    stat_sheet['C134'].value = 'Test successful coverage'
    stat_sheet['C134'].font = bold_font
    stat_sheet['E134'].value = '=D131/(H131-G131)'
    stat_sheet['E134'].font = bold_font
    stat_sheet['E134'].number_format = '0.00%'

    print("   Updated 'Test Statistics' formulas to cover rows 11 to 130 (120 UCs total)!")

    # 6. Update Cover sheet
    print("\n📜 6. Updating Cover sheet...")
    cover_sheet = wb['Cover']
    cover_sheet['F6'].value = 'v1.1'
    # Add record of change at row 12
    row_change = 12
    cover_sheet[f'A{row_change}'].value = '01/10/2026'
    cover_sheet[f'B{row_change}'].value = 'v1.1'
    cover_sheet[f'C{row_change}'].value = 'Test Cases & Statistics (UC-66 to UC-120)'
    cover_sheet[f'D{row_change}'].value = 'A'
    cover_sheet[f'E{row_change}'].value = 'Extended full enterprise test suite with UC-66 to UC-120 covering Sprint 5 to Sprint 8'
    cover_sheet[f'F{row_change}'].value = 'Report 4 – Software Design Document (SDD v2.0)'

    # 7. Save workbook locally
    print(f"\n💾 7. Saving workbook to {LOCAL_OUTPUT}...")
    os.makedirs('docs', exist_ok=True)
    wb.save(LOCAL_OUTPUT)
    wb.save('/tmp/Report5_Test_Report_Updated.xlsx')
    print(f"✅ Saved successfully ({os.path.getsize(LOCAL_OUTPUT)} bytes)!")

    # 8. Upload directly to Google Drive
    print(f"\n☁️ 8. Directly overwriting Google Drive file: {SPREADSHEET_ID}...")
    token = get_access_token(SA_PATH)
    upload_url = f"https://www.googleapis.com/upload/drive/v3/files/{SPREADSHEET_ID}?uploadType=media"

    with open(LOCAL_OUTPUT, 'rb') as f:
        file_bytes = f.read()

    req = urllib.request.Request(upload_url, data=file_bytes, headers={
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    }, method='PATCH')

    with urllib.request.urlopen(req) as resp:
        result = json.loads(resp.read().decode('utf-8'))
        print("\n" + "="*60)
        print("🎉 GHI ĐÈ THÀNH CÔNG 100% TRỰC TIẾP LÊN LINK GOOGLE SPREADSHEET!")
        print(f"📄 Tên file trên Cloud: {result.get('name')}")
        print(f"🆔 File ID: {result.get('id')}")
        print(f"🔗 Link Google Sheet: https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/edit")
        print("="*60 + "\n")

if __name__ == '__main__':
    update_test_report()
